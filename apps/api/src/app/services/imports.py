"""Bounded import parsing and owner-approved Timecue import execution.

The service keeps spreadsheet rows in an IntelliQ-owned ledger before any
upstream mutation is possible.  External keys and row keys remain stable when
the same source is previewed again; Timecue IDs are mappings learned from
explicit reads or from an earlier committed ledger.  Only the small set of
routes documented in ``docs/v4-timecue-capabilities.md`` is executable here.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Protocol
from uuid import uuid4
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from src.app.integrations.timecue import UpstreamIntegrationError, _public_error_detail
from src.app.persistence.store import SessionRecord, Store

MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_IMPORT_FILES = 16
MAX_IMPORT_ROWS = 5_000
MAX_CELL_CHARS = 16_000
MAX_XLSX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_XLSX_MEMBERS = 512
IMPORT_RECORD_KIND = "import_batches"
IMPORT_COMMIT_LEASE = "imports:commit"
SCHEMA_VERSION = 1
DEFAULT_PROJECT_CURRENCY = "PLN"
DEFAULT_PROJECT_COUNTRY_CODE = "PL"
TIMECUE_PROJECT_STATUSES = frozenset(
    {
        "lead",
        "estimation",
        "waiting_for_client",
        "accepted",
        "scheduled",
        "in_progress",
        "done",
        "archived",
    }
)
PROJECT_STATUS_ALIASES = {
    "planned": "scheduled",
    "planning": "scheduled",
    "active": "in_progress",
    "completed": "done",
    "complete": "done",
    "cancelled": "archived",
    "canceled": "archived",
}

ENTITY_ORDER = ("projects", "tasks", "workers", "skills", "assignments", "planning")
ENTITY_SHEETS = {
    "projects": "Projects",
    "tasks": "Tasks",
    "workers": "Workers",
    "skills": "Skills",
    "assignments": "Assignments",
    "planning": "Planning",
}
ENTITY_ALIASES = {
    "project": "projects",
    "projects": "projects",
    "task": "tasks",
    "tasks": "tasks",
    "worker": "workers",
    "workers": "workers",
    "member": "workers",
    "members": "workers",
    "skill": "skills",
    "skills": "skills",
    "assignment": "assignments",
    "assignments": "assignments",
    "planning": "planning",
    "plan": "planning",
}
REQUIRED_COLUMNS: dict[str, tuple[str, ...]] = {
    "projects": ("externalKey", "name", "timezone", "address", "targetFinishAt"),
    "tasks": ("externalKey", "projectKey", "title", "plannedStart", "plannedEnd"),
    "workers": ("externalKey", "timecueMemberId", "skillKeys"),
    "skills": ("externalKey", "name"),
    "assignments": ("taskKey", "workerKey"),
    "planning": (),
}
TEMPLATE_COLUMNS: dict[str, tuple[str, ...]] = {
    "projects": (
        "externalKey",
        "name",
        "timezone",
        "address",
        "defaultLatitude",
        "defaultLongitude",
        "targetFinishAt",
        "priority",
        "description",
        "status",
        "currency",
        "countryCode",
        "timecueProjectId",
    ),
    "tasks": (
        "externalKey",
        "projectKey",
        "title",
        "plannedStart",
        "plannedEnd",
        "description",
        "status",
        "estimatedMinutes",
        "workType",
        "weatherRulesJson",
        "weatherCoverage",
        "timecueTaskId",
    ),
    "workers": ("externalKey", "timecueMemberId", "skillKeys"),
    "skills": ("externalKey", "name", "description", "timecueSpecialtyId"),
    "assignments": ("taskKey", "workerKey"),
    "planning": (
        "assumptionType",
        "taskKey",
        "projectKey",
        "fromProjectKey",
        "toProjectKey",
        "workerKey",
        "optimisticPersonHours",
        "mostLikelyPersonHours",
        "pessimisticPersonHours",
        "prerequisiteKeys",
        "specialtyKeys",
        "minCrew",
        "maxCrew",
        "target",
        "priority",
        "outboundTravelHours",
        "returnTravelHours",
        "setupHours",
        "startsAt",
        "endsAt",
        "transferConfirmed",
    ),
}
CSV_FILENAMES = {entity: f"{entity}.csv" for entity in ENTITY_SHEETS}


class ImportServiceError(ValueError):
    """An actionable import validation, persistence, or upstream error."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        status_code: int = 422,
        details: Sequence[Mapping[str, object]] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code
        self.details = [dict(item) for item in details or ()]


@dataclass(frozen=True, slots=True)
class UploadedImport:
    """A bounded upload copied out of Starlette before parsing."""

    filename: str
    content: bytes
    content_type: str | None = None


class RouteClient(Protocol):
    """The narrow upstream seam used by commit and injectable in tests."""

    def request(
        self,
        method: str,
        path: str,
        *,
        params: Mapping[str, str | int] | None = None,
        json_body: Mapping[str, object] | None = None,
    ) -> object:
        """Perform one authenticated existing-route request."""


class TimecueRouteClient:
    """Use the live adapter's per-session request boundary for import commits.

    The live adapter intentionally exposes reads rather than a shared mutable
    HTTP client.  This local writer seam delegates to its short-lived request
    method, refreshes once after a 401, and never retries an ambiguous request.
    """

    def __init__(self, adapter: object, session: SessionRecord) -> None:
        self.adapter = adapter
        self.session = session

    def request(
        self,
        method: str,
        path: str,
        *,
        params: Mapping[str, str | int] | None = None,
        json_body: Mapping[str, object] | None = None,
    ) -> object:
        adapter_request = getattr(self.adapter, "_request", None)
        adapter_tokens = getattr(self.adapter, "_tokens", None)
        adapter_refresh = getattr(self.adapter, "refresh", None)
        if callable(adapter_request) and callable(adapter_tokens):
            try:
                failed_access = str(adapter_tokens(self.session).access)
            except Exception as exc:  # pragma: no cover - adapter boundary guard
                raise ImportServiceError(
                    "timecue_session_unavailable",
                    "The authenticated Timecue session is unavailable.",
                    status_code=502,
                ) from exc
            kwargs: dict[str, object] = {}
            if params is not None:
                kwargs["params"] = dict(params)
            if json_body is not None:
                kwargs["json"] = dict(json_body)
            try:
                response = adapter_request(self.session, method, path, **kwargs)
            except UpstreamIntegrationError as exc:
                if exc.status_code != 401 or not callable(adapter_refresh):
                    raise ImportServiceError(
                        "timecue_request_failed",
                        _upstream_rejection_message(exc),
                        status_code=502,
                    ) from exc
                try:
                    adapter_refresh(self.session, failed_access)
                    response = adapter_request(self.session, method, path, **kwargs)
                except UpstreamIntegrationError as retry_exc:
                    raise ImportServiceError(
                        "timecue_request_failed",
                        _upstream_rejection_message(retry_exc, after_refresh=True),
                        status_code=502,
                    ) from retry_exc
            return _response_payload(response)

        public_request = getattr(self.adapter, "request", None)
        if callable(public_request):
            try:
                return public_request(method, path, params=params, json_body=json_body)
            except TypeError:
                # Small recording fakes often name the body argument ``json``.
                return public_request(method, path, params=params, json=json_body)
        raise ImportServiceError(
            "timecue_write_unsupported",
            "This Timecue adapter does not expose the existing-route request seam.",
            status_code=502,
        )


def preview_import(
    store: Store,
    organization_id: str,
    uploads: Sequence[UploadedImport],
    *,
    member_rows: Sequence[Mapping[str, object]] = (),
    source_revision: str | None = None,
    current_snapshot: Mapping[str, object] | None = None,
) -> dict[str, object]:
    """Parse, validate and persist a reviewable import without upstream writes."""

    if not uploads:
        raise ImportServiceError(
            "import_file_required", "At least one CSV or XLSX file is required."
        )
    total_bytes = sum(len(upload.content) for upload in uploads)
    if total_bytes > MAX_IMPORT_BYTES:
        raise ImportServiceError(
            "import_file_too_large",
            f"Import uploads must not exceed {MAX_IMPORT_BYTES // (1024 * 1024)} MiB.",
        )

    tables: list[ParsedTable] = []
    parse_errors: list[dict[str, object]] = []
    for upload in uploads:
        try:
            parsed_tables, errors = _parse_upload(upload)
        except ImportServiceError as exc:
            parse_errors.append(
                {"code": exc.code, "message": exc.message, "file": _safe_filename(upload.filename)}
            )
            continue
        tables.extend(parsed_tables)
        parse_errors.extend(errors)

    row_count = sum(len(table.rows) for table in tables)
    if row_count > MAX_IMPORT_ROWS:
        parse_errors.append(
            {
                "code": "import_row_limit_exceeded",
                "message": f"The import is limited to {MAX_IMPORT_ROWS} data rows.",
            }
        )

    rows = _build_row_records(tables)
    previous_mappings = _previous_mappings(store.list_records(organization_id, IMPORT_RECORD_KIND))
    for row in rows:
        mapping = previous_mappings.get(_mapping_key(row))
        if mapping:
            row["mapping"] = dict(mapping)

    _validate_rows(rows, member_rows=member_rows, current_snapshot=current_snapshot)
    errors = [*parse_errors, *_row_errors(rows)]
    counts = _counts(rows, errors)
    batch_id = str(uuid4())
    source_files = _source_files(uploads, tables)
    record: dict[str, object] = {
        "schemaVersion": SCHEMA_VERSION,
        "id": batch_id,
        "organizationId": organization_id,
        "status": "needs_correction" if errors else "ready",
        "sourceRevision": source_revision,
        "sourceFiles": source_files,
        "sourceHash": _source_hash(source_files),
        "rows": rows,
        "errors": errors,
        "counts": counts,
        "planningStatus": "proposed"
        if any(
            row["entity"] == "planning"
            or any(
                _row_values(row).get(field) not in (None, "")
                for field in (
                    "optimisticPersonHours",
                    "mostLikelyPersonHours",
                    "pessimisticPersonHours",
                    "prerequisiteKeys",
                    "specialtyKeys",
                    "skillKeys",
                    "targetFinishAt",
                    "priority",
                    "minCrew",
                    "maxCrew",
                    "workType",
                    "weatherRulesJson",
                )
            )
            for row in rows
        )
        else "none",
        "receipts": [],
        "message": (
            "Correct the listed rows before committing this import."
            if errors
            else "Preview is ready for explicit owner confirmation."
        ),
    }
    record["previewHash"] = _record_hash(record)
    return store.save_record(organization_id, IMPORT_RECORD_KIND, batch_id, record)


def commit_import(
    store: Store,
    organization_id: str,
    import_id: str,
    *,
    upstream: RouteClient | None,
    actor_id: str | None = None,
    confirm_planning: bool = False,
    expected_preview_hash: str | None = None,
    resume: bool = False,
) -> dict[str, object]:
    """Serialize one import commit per organization with a durable lease."""

    lease = store.acquire_lease(organization_id, IMPORT_COMMIT_LEASE)
    if lease is None:
        raise ImportServiceError(
            "import_commit_in_progress",
            "Another import commit is already running for this organization.",
            status_code=409,
        )
    try:
        return _commit_import_locked(
            store,
            organization_id,
            import_id,
            upstream=upstream,
            actor_id=actor_id,
            confirm_planning=confirm_planning,
            expected_preview_hash=expected_preview_hash,
            resume=resume,
        )
    finally:
        store.release_lease(organization_id, IMPORT_COMMIT_LEASE, lease)


def _commit_import_locked(
    store: Store,
    organization_id: str,
    import_id: str,
    *,
    upstream: RouteClient | None,
    actor_id: str | None = None,
    confirm_planning: bool = False,
    expected_preview_hash: str | None = None,
    resume: bool = False,
) -> dict[str, object]:
    """Apply supported rows with per-step receipts and canonical readbacks.

    Network calls are intentionally outside the local record write.  Each step
    is persisted immediately after its outcome, so a process restart can show
    the partial state and an explicit resume can skip receipts already applied.
    """

    record = store.get_record(organization_id, IMPORT_RECORD_KIND, import_id)
    if record is None:
        raise ImportServiceError(
            "import_not_found", "The import batch was not found.", status_code=404
        )
    if expected_preview_hash and record.get("previewHash") != expected_preview_hash:
        raise ImportServiceError(
            "import_preview_changed",
            "The import preview changed; review it again before committing.",
            status_code=409,
        )
    if record.get("errors"):
        raise ImportServiceError(
            "import_requires_correction",
            "The import has validation errors and cannot be committed.",
            status_code=409,
            details=_mapping_list(record.get("errors")),
        )
    if record.get("status") == "applied":
        return record
    if record.get("status") in {"partially_applied", "needs_reconciliation"} and not resume:
        return {
            **record,
            "message": (
                "This batch has existing receipts; request an explicit resume after reconciliation."
            ),
        }

    rows = [_mutable_mapping(item) for item in _mapping_list(record.get("rows"))]
    receipts = [_mutable_mapping(item) for item in _mapping_list(record.get("receipts"))]
    receipt_by_row = {str(item.get("rowKey")): item for item in receipts if item.get("rowKey")}
    record["rows"] = rows
    record["receipts"] = receipts
    if confirm_planning and record.get("planningStatus") == "proposed":
        record["planningConfirmationRequestedBy"] = actor_id
        record["planningConfirmationRequestedAt"] = _utc_now()

    if upstream is None:
        for row in rows:
            if str(row.get("entity")) not in {"projects", "tasks", "assignments"}:
                continue
            row_key = str(row.get("rowKey"))
            if _receipt_is_terminal(receipt_by_row.get(row_key)):
                continue
            receipt = _receipt(
                row,
                status="manual",
                code="live_timecue_required",
                message="Fixture mode does not write to Timecue; perform this step manually.",
            )
            _upsert_receipt(receipts, receipt_by_row, receipt)
        record["message"] = "No upstream write was attempted; live Timecue access is required."
        record["status"] = "needs_reconciliation"
        return store.save_record(organization_id, IMPORT_RECORD_KIND, import_id, record)

    project_ids: dict[str, str] = {}
    task_ids: dict[str, str] = {}
    worker_profile_ids: dict[str, str] = {}

    _seed_existing_mappings(rows, project_ids, task_ids, worker_profile_ids)
    _commit_projects(
        rows,
        project_ids,
        upstream,
        receipts,
        receipt_by_row,
        organization_id,
        store,
        record,
    )
    _commit_tasks(
        rows,
        project_ids,
        task_ids,
        upstream,
        receipts,
        receipt_by_row,
        organization_id,
        store,
        record,
    )
    _commit_assignments(
        rows,
        project_ids,
        task_ids,
        worker_profile_ids,
        upstream,
        receipts,
        receipt_by_row,
        organization_id,
        store,
        record,
    )

    for row in rows:
        entity = str(row.get("entity"))
        if entity in {"projects", "tasks", "assignments"}:
            continue
        row_key = str(row.get("rowKey"))
        if _receipt_is_terminal(receipt_by_row.get(row_key)):
            continue
        receipt = _receipt(
            row, status="verified", message="Validated in IntelliQ; no upstream write required."
        )
        _upsert_receipt(receipts, receipt_by_row, receipt)

    record["rows"] = rows
    record["receipts"] = receipts
    record["status"] = _batch_status(receipts)
    record["message"] = _batch_message(record["status"])
    if confirm_planning and record.get("planningStatus") == "proposed":
        if _planning_merge_ready(rows, receipts):
            try:
                merged = merge_confirmed_planning(store, organization_id, rows)
            except ImportServiceError as exc:
                record["planningStatus"] = "needs_reconciliation"
                record["planningMerge"] = {
                    "status": "blocked",
                    "code": exc.code,
                    "message": exc.message,
                }
                record["status"] = "needs_reconciliation"
                record["message"] = (
                    "Supported upstream rows were read back, but planning needs reconciliation."
                )
            else:
                record["planningStatus"] = "confirmed"
                record["planningConfirmedBy"] = actor_id
                record["planningConfirmedAt"] = _utc_now()
                record["planningVersion"] = merged.get("version")
                record["planningMerge"] = {
                    "status": "applied",
                    "version": merged.get("version"),
                    "source": "confirmed_import",
                }
        else:
            record["planningMerge"] = {
                "status": "blocked",
                "code": "upstream_receipts_incomplete",
                "message": (
                    "Planning waits for supported upstream rows to be applied and read back."
                ),
            }
    return store.save_record(organization_id, IMPORT_RECORD_KIND, import_id, record)


def merge_confirmed_planning(
    store: Store,
    organization_id: str,
    rows: Sequence[Mapping[str, object]],
) -> dict[str, object]:
    """Merge verified import assumptions into the versioned planning overlay.

    The merge is additive and ID-scoped: it updates only entities represented by
    the import, keeps unrelated manager assumptions, and never removes an
    upstream entity.  Callers must invoke this only after supported Timecue
    mappings have been read back and the owner has explicitly confirmed the
    planning portion of the import.
    """

    saved = store.get_planning(organization_id) or {}
    current = dict(saved)
    try:
        expected_version = int(current.get("version", 0) or 0)
    except (TypeError, ValueError) as exc:
        raise ImportServiceError(
            "planning_version_invalid",
            "The existing planning overlay has an invalid version.",
            status_code=409,
        ) from exc
    projects = _mapping_list(current.get("projects"))
    tasks = _mapping_list(current.get("tasks"))
    workers = _mapping_list(current.get("workers"))
    transfers = _mapping_list(current.get("transfers"))

    project_ids = _import_mapping_ids(rows, "projects", "upstreamId")
    task_ids = _import_mapping_ids(rows, "tasks", "upstreamId")
    worker_ids = _import_mapping_ids(rows, "workers", "workerProfileId")
    skill_ids = _import_mapping_ids(rows, "skills", "upstreamId")

    project_updates: list[dict[str, object]] = []
    for row in rows:
        if row.get("entity") != "projects":
            continue
        values = _row_values(row)
        external_key = _string_value(values.get("externalKey"))
        project_id = project_ids.get(external_key or "")
        target_finish = _string_value(values.get("targetFinishAt"))
        timezone = _string_value(values.get("timezone"))
        if not project_id or not target_finish or not timezone:
            raise ImportServiceError(
                "planning_project_mapping_missing",
                (
                    f"Project {external_key or row.get('rowKey')} is not fully mapped "
                    "or lacks targetFinishAt/timezone."
                ),
                status_code=409,
            )
        update: dict[str, object] = {
            "id": project_id,
            "targetFinishAt": target_finish,
            "timezone": timezone,
        }
        priority = _number_value(values.get("priority"))
        if priority is not None:
            update["priority"] = int(priority) if priority.is_integer() else priority
        latitude = _number_value(values.get("defaultLatitude"))
        longitude = _number_value(values.get("defaultLongitude"))
        if latitude is not None and longitude is not None:
            update["location"] = {"latitude": latitude, "longitude": longitude}
        project_updates.append(update)

    planning_by_task = {
        _string_value(_row_values(row).get("taskKey")): _row_values(row)
        for row in rows
        if row.get("entity") == "planning"
        and not _is_transfer_row(_row_values(row))
        and _string_value(_row_values(row).get("taskKey"))
    }
    task_updates: list[dict[str, object]] = []
    for row in rows:
        if row.get("entity") != "tasks":
            continue
        values = _row_values(row)
        external_key = _string_value(values.get("externalKey"))
        task_id = task_ids.get(external_key or "")
        project_key = _string_value(values.get("projectKey"))
        project_id = project_ids.get(project_key or "")
        if not task_id or not project_id:
            raise ImportServiceError(
                "planning_task_mapping_missing",
                f"Task {external_key or row.get('rowKey')} is not fully mapped to a project.",
                status_code=409,
            )
        planning_values = planning_by_task.get(external_key or "", {})

        update: dict[str, object] = {"id": task_id, "projectId": project_id}
        effort = {
            name: _number_value(_assumption_value(planning_values, values, field))
            for name, field in (
                ("optimistic", "optimisticPersonHours"),
                ("mostLikely", "mostLikelyPersonHours"),
                ("pessimistic", "pessimisticPersonHours"),
            )
        }
        if all(value is not None for value in effort.values()):
            update["remainingPersonHours"] = {
                name: float(value) for name, value in effort.items() if value is not None
            }
        prerequisites = [
            task_ids[key]
            for key in _split_keys(_assumption_value(planning_values, values, "prerequisiteKeys"))
            if key in task_ids
        ]
        if prerequisites:
            update["predecessorIds"] = prerequisites
        specialties = [
            skill_ids[key]
            for key in _split_keys(_assumption_value(planning_values, values, "specialtyKeys"))
            if key in skill_ids
        ]
        if specialties:
            update["requiredSpecialtyId"] = specialties[0]
        for source_key, target_key in (
            ("minCrew", "minCrew"),
            ("maxCrew", "maxCrew"),
            ("priority", "priority"),
        ):
            value = _number_value(_assumption_value(planning_values, values, source_key))
            if value is not None:
                update[target_key] = int(value) if value.is_integer() else value
        work_type = _string_value(values.get("workType"))
        if work_type:
            update["workType"] = work_type
        weather_rules = _json_object(values.get("weatherRulesJson"))
        if _confirmed_weather_rules(weather_rules):
            update["weatherRules"] = weather_rules
        task_updates.append(update)

    worker_updates: list[dict[str, object]] = []
    for row in rows:
        if row.get("entity") != "workers":
            continue
        values = _row_values(row)
        external_key = _string_value(values.get("externalKey"))
        worker_id = worker_ids.get(external_key or "")
        if not worker_id:
            continue
        update: dict[str, object] = {"id": worker_id}
        specialties = [
            skill_ids[key] for key in _split_keys(values.get("skillKeys")) if key in skill_ids
        ]
        if specialties:
            update["specialtyIds"] = specialties
        worker_updates.append(update)

    transfer_updates: list[dict[str, object]] = []
    for row in rows:
        if row.get("entity") != "planning" or not _is_transfer_row(_row_values(row)):
            continue
        values = _row_values(row)
        from_project = project_ids.get(_string_value(values.get("fromProjectKey")) or "")
        to_project = project_ids.get(_string_value(values.get("toProjectKey")) or "")
        worker_id = worker_ids.get(_string_value(values.get("workerKey")) or "")
        if not from_project or not to_project or not worker_id:
            raise ImportServiceError(
                "planning_transfer_mapping_missing",
                f"Transfer {row.get('rowKey')} references an unmapped project or worker.",
                status_code=409,
            )
        transfer: dict[str, object] = {
            "id": f"import:{row.get('rowKey')}",
            "fromProjectId": from_project,
            "toProjectId": to_project,
            "workerId": worker_id,
            "confirmed": _boolean_value(values.get("transferConfirmed"), default=True),
        }
        for field in ("outboundTravelHours", "returnTravelHours", "setupHours"):
            value = _number_value(values.get(field))
            if value is not None:
                transfer[field] = value
        for field in ("startsAt", "endsAt"):
            value = _string_value(values.get(field))
            if value:
                transfer[field] = value
        transfer_updates.append(transfer)

    merged = {
        **current,
        "projects": _merge_overlay_items(projects, project_updates),
        "tasks": _merge_overlay_items(tasks, task_updates),
        "workers": _merge_overlay_items(workers, worker_updates),
        "transfers": _merge_overlay_items(transfers, transfer_updates),
    }
    try:
        return store.save_planning(organization_id, merged, expected_version)
    except ValueError as exc:
        raise ImportServiceError(
            "planning_version_conflict",
            "Planning changed while this import was being confirmed; review and retry.",
            status_code=409,
        ) from exc


def _planning_merge_ready(
    rows: Sequence[Mapping[str, object]], receipts: Sequence[Mapping[str, object]]
) -> bool:
    receipt_by_row = {str(item.get("rowKey")): item for item in receipts}
    supported = {"projects", "tasks", "assignments"}
    return all(
        receipt_by_row.get(str(row.get("rowKey")), {}).get("status") in {"applied", "verified"}
        for row in rows
        if row.get("entity") in supported
    ) and all(
        receipt_by_row.get(str(row.get("rowKey")), {}).get("status") in {"applied", "verified"}
        for row in rows
    )


def _import_mapping_ids(
    rows: Sequence[Mapping[str, object]], entity: str, mapping_key: str
) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in rows:
        if row.get("entity") != entity:
            continue
        external_key = _string_value(_row_values(row).get("externalKey"))
        mapped_id = _string_value(_mapping(row.get("mapping")).get(mapping_key))
        if external_key and mapped_id:
            result[external_key] = mapped_id
    return result


def _merge_overlay_items(
    existing: Sequence[Mapping[str, object]], updates: Sequence[Mapping[str, object]]
) -> list[dict[str, object]]:
    by_id = {str(item.get("id")): dict(item) for item in existing if item.get("id")}
    order = [str(item.get("id")) for item in existing if item.get("id")]
    for item in updates:
        identifier = _string_value(item.get("id"))
        if not identifier:
            continue
        if identifier not in by_id:
            order.append(identifier)
        by_id[identifier] = {**by_id.get(identifier, {}), **dict(item)}
    return [by_id[identifier] for identifier in order if identifier in by_id]


def _is_transfer_row(values: Mapping[str, object]) -> bool:
    assumption_type = (_string_value(values.get("assumptionType")) or "task").lower()
    return assumption_type in {"transfer", "transfer_assumption", "transferassumption"}


def _json_object(value: object) -> dict[str, object] | None:
    text = _string_value(value)
    if not text:
        return None
    try:
        decoded = json.loads(text)
    except json.JSONDecodeError:
        return None
    return dict(decoded) if isinstance(decoded, Mapping) else None


def _confirmed_weather_rules(value: Mapping[str, object] | None) -> bool:
    if value is None or value.get("confirmed") is not True:
        return False
    rules = value.get("rules")
    if not isinstance(rules, list) or not rules:
        return False
    capacity = value.get("capacity", value.get("defaultCapacity"))
    return (
        isinstance(capacity, (int, float)) and not isinstance(capacity, bool) and 0 <= capacity <= 1
    )


def _boolean_value(value: object, *, default: bool) -> bool:
    if isinstance(value, bool):
        return value
    text = _string_value(value)
    if text is None:
        return default
    if text.lower() in {"true", "yes", "1", "confirmed"}:
        return True
    if text.lower() in {"false", "no", "0", "unconfirmed"}:
        return False
    return default


def template_bytes(format_name: str) -> tuple[bytes, str, str]:
    """Return a deterministic XLSX or ZIP-of-CSV template download."""

    normalized = format_name.strip().lower()
    if normalized == "csv":
        import zipfile as csv_zip

        output = io.BytesIO()
        with csv_zip.ZipFile(output, "w", compression=csv_zip.ZIP_DEFLATED) as archive:
            for entity in ENTITY_ORDER:
                body = ",".join(TEMPLATE_COLUMNS[entity]) + "\n"
                archive.writestr(CSV_FILENAMES[entity], body.encode("utf-8"))
        return output.getvalue(), "application/zip", "intelliq-import-template-csv.zip"
    if normalized != "xlsx":
        raise ImportServiceError("template_format_invalid", "Template format must be xlsx or csv.")
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - runtime dependency boundary
        raise ImportServiceError(
            "xlsx_parser_unavailable",
            "XLSX support requires the openpyxl runtime dependency.",
            status_code=503,
        ) from exc
    workbook = Workbook()
    first = workbook.active
    first.title = ENTITY_SHEETS[ENTITY_ORDER[0]]
    for entity in ENTITY_ORDER:
        sheet = first if entity == ENTITY_ORDER[0] else workbook.create_sheet(ENTITY_SHEETS[entity])
        sheet.append(list(TEMPLATE_COLUMNS[entity]))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{_column_letter(len(TEMPLATE_COLUMNS[entity]))}1"
    output = io.BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "intelliq-import-template.xlsx",
    )


@dataclass(slots=True)
class ParsedTable:
    entity: str
    source_file: str
    sheet: str
    rows: list[tuple[int, dict[str, object]]]


def _parse_upload(upload: UploadedImport) -> tuple[list[ParsedTable], list[dict[str, object]]]:
    filename = _safe_filename(upload.filename)
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _parse_csv(upload.content, filename)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _parse_xlsx(upload.content, filename)
    raise ImportServiceError(
        "import_file_type_unsupported",
        "Only CSV and XLSX files are accepted; macros and legacy XLS files are not supported.",
    )


def _parse_csv(content: bytes, filename: str) -> tuple[list[ParsedTable], list[dict[str, object]]]:
    entity = _entity_from_filename(filename)
    if entity is None:
        raise ImportServiceError(
            "csv_entity_unknown",
            (
                "CSV names must be projects.csv, tasks.csv, workers.csv, skills.csv, "
                "assignments.csv, or planning.csv."
            ),
        )
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportServiceError("csv_encoding_invalid", "CSV files must use UTF-8.") from exc
    if "\x00" in text:
        raise ImportServiceError("csv_content_invalid", "CSV contains a NUL byte.")
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        raw_headers = next(reader)
    except StopIteration as exc:
        raise ImportServiceError("csv_header_missing", "CSV must contain a header row.") from exc
    headers, header_errors = _canonical_headers(raw_headers)
    errors = [
        {"code": code, "message": message, "file": filename} for code, message in header_errors
    ]
    rows: list[tuple[int, dict[str, object]]] = []
    for row_number, raw_values in enumerate(reader, start=2):
        if not any(str(value).strip() for value in raw_values):
            continue
        values = _row_from_values(headers, raw_values)
        if any(_is_formula(value) for value in values.values()):
            errors.append(
                {
                    "code": "formula_not_allowed",
                    "message": "Spreadsheet formulas are not accepted.",
                    "file": filename,
                    "row": row_number,
                }
            )
        rows.append((row_number, values))
        if len(rows) > MAX_IMPORT_ROWS:
            break
    return [ParsedTable(entity, filename, ENTITY_SHEETS[entity], rows)], errors


def _parse_xlsx(content: bytes, filename: str) -> tuple[list[ParsedTable], list[dict[str, object]]]:
    if not content.startswith(b"PK"):
        raise ImportServiceError(
            "xlsx_content_invalid", "The XLSX file is not a valid ZIP workbook."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_MEMBERS:
                raise ImportServiceError(
                    "xlsx_member_limit_exceeded", "The workbook contains too many ZIP members."
                )
            if sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ImportServiceError(
                    "xlsx_decompressed_limit_exceeded",
                    "The workbook's decompressed size exceeds the safety bound.",
                )
            names = {info.filename.lower() for info in infos}
            if any(
                "vbaproject.bin" in name or "externallinks/" in name or "external-links" in name
                for name in names
            ):
                raise ImportServiceError(
                    "xlsx_external_content_not_allowed",
                    "Macros and external workbook links are not accepted.",
                )
    except zipfile.BadZipFile as exc:
        raise ImportServiceError(
            "xlsx_content_invalid", "The XLSX file is not a valid ZIP workbook."
        ) from exc
    if filename.lower().endswith((".xlsm", ".xltm")):
        raise ImportServiceError(
            "xlsx_macros_not_allowed", "Macro-enabled workbooks are not accepted."
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - runtime dependency boundary
        raise ImportServiceError(
            "xlsx_parser_unavailable",
            "XLSX support requires the openpyxl runtime dependency.",
            status_code=503,
        ) from exc
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=True
        )
    except (ValueError, OSError, zipfile.BadZipFile) as exc:
        raise ImportServiceError(
            "xlsx_content_invalid", "The XLSX workbook could not be read."
        ) from exc
    tables: list[ParsedTable] = []
    errors: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        entity = _entity_from_sheet(sheet.title)
        if entity is None:
            errors.append(
                {
                    "code": "xlsx_sheet_unknown",
                    "message": f"Sheet {sheet.title!r} is not part of the import template.",
                    "file": filename,
                }
            )
            continue
        iterator = sheet.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration:
            errors.append(
                {
                    "code": "xlsx_header_missing",
                    "message": f"Sheet {sheet.title!r} has no header row.",
                    "file": filename,
                }
            )
            continue
        raw_headers = [_cell_value(cell) for cell in header_cells]
        headers, header_errors = _canonical_headers(raw_headers)
        errors.extend(
            {"code": code, "message": message, "file": filename, "sheet": sheet.title}
            for code, message in header_errors
        )
        rows: list[tuple[int, dict[str, object]]] = []
        for row_number, cells in enumerate(iterator, start=2):
            raw_values = [_cell_value(cell) for cell in cells]
            if not any(_string_value(value) for value in raw_values):
                continue
            if any(_cell_formula(cell) for cell in cells):
                errors.append(
                    {
                        "code": "formula_not_allowed",
                        "message": "Spreadsheet formulas are not accepted.",
                        "file": filename,
                        "sheet": sheet.title,
                        "row": row_number,
                    }
                )
            rows.append((row_number, _row_from_values(headers, raw_values)))
            if sum(len(table.rows) for table in tables) + len(rows) > MAX_IMPORT_ROWS:
                break
        tables.append(ParsedTable(entity, filename, sheet.title, rows))
    workbook.close()
    return tables, errors


def _build_row_records(tables: Sequence[ParsedTable]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for table in tables:
        for source_row, values in table.rows:
            row_key = _stable_row_key(table.entity, values, source_row)
            rows.append(
                {
                    "entity": table.entity,
                    "sheet": table.sheet,
                    "sourceFile": table.source_file,
                    "sourceRow": source_row,
                    "rowKey": row_key,
                    "key": row_key,
                    "stableRowKey": row_key,
                    "values": values,
                    "status": "valid",
                    "errors": [],
                    "mapping": {},
                    "operation": _operation_for(table.entity),
                }
            )
    return rows


def _validate_rows(
    rows: list[dict[str, object]],
    *,
    member_rows: Sequence[Mapping[str, object]],
    current_snapshot: Mapping[str, object] | None,
) -> None:
    by_entity_external: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        entity = str(row.get("entity"))
        values = _row_values(row)
        errors = _string_list(row.get("errors"))
        for column in REQUIRED_COLUMNS.get(entity, ()):
            if not _string_value(values.get(column)):
                errors.append(f"missing_required:{column}")
        if entity == "planning":
            if _is_transfer_row(values):
                for column in ("fromProjectKey", "toProjectKey", "workerKey"):
                    if not _string_value(values.get(column)):
                        errors.append(f"missing_required:{column}")
            elif not _string_value(values.get("taskKey")):
                errors.append("missing_required:taskKey")
        external_key = _string_value(values.get("externalKey"))
        if external_key and entity != "assignments":
            by_entity_external[(entity, external_key)].append(row)
        if entity == "projects":
            timezone_value = _string_value(values.get("timezone"))
            if timezone_value:
                try:
                    ZoneInfo(timezone_value)
                except ZoneInfoNotFoundError:
                    errors.append("timezone_invalid")
        elif entity == "tasks":
            _validate_datetime(values.get("plannedStart"), "plannedStart", errors)
            _validate_datetime(values.get("plannedEnd"), "plannedEnd", errors)
            _validate_ordered_datetimes(
                values.get("plannedStart"), values.get("plannedEnd"), errors
            )
            _validate_coordinates(values, errors)
            _validate_weather_rules(values, errors)
        elif entity == "planning":
            _validate_planning_values(values, errors)
            _validate_datetime(values.get("startsAt"), "startsAt", errors)
            _validate_datetime(values.get("endsAt"), "endsAt", errors)
            _validate_ordered_datetimes(values.get("startsAt"), values.get("endsAt"), errors)
        row["errors"] = errors

    for (entity, external_key), duplicate_rows in by_entity_external.items():
        if len(duplicate_rows) < 2:
            continue
        for row in duplicate_rows:
            _append_error(row, f"duplicate_external_key:{entity}:{external_key}")

    keys = {
        entity: {
            _string_value(_row_values(row).get("externalKey"))
            for row in rows
            if row.get("entity") == entity and _string_value(_row_values(row).get("externalKey"))
        }
        for entity in ENTITY_SHEETS
    }
    for row in rows:
        entity = str(row.get("entity"))
        values = _row_values(row)
        if entity == "tasks":
            _require_reference(row, "projectKey", values.get("projectKey"), keys["projects"])
        elif entity == "workers":
            _match_member(row, member_rows)
            for skill_key in _split_keys(values.get("skillKeys")):
                _require_reference(row, "skillKeys", skill_key, keys["skills"])
        elif entity == "assignments":
            _require_reference(row, "taskKey", values.get("taskKey"), keys["tasks"])
            _require_reference(row, "workerKey", values.get("workerKey"), keys["workers"])
        elif entity == "planning":
            if _is_transfer_row(values):
                _require_reference(
                    row, "fromProjectKey", values.get("fromProjectKey"), keys["projects"]
                )
                _require_reference(
                    row, "toProjectKey", values.get("toProjectKey"), keys["projects"]
                )
                _require_reference(row, "workerKey", values.get("workerKey"), keys["workers"])
            else:
                _require_reference(row, "taskKey", values.get("taskKey"), keys["tasks"])
            for prerequisite in _split_keys(values.get("prerequisiteKeys")):
                _require_reference(row, "prerequisiteKeys", prerequisite, keys["tasks"])
            for specialty in _split_keys(values.get("specialtyKeys")):
                _require_reference(row, "specialtyKeys", specialty, keys["skills"])
        _validate_explicit_upstream_id(row, current_snapshot)

    _mark_planning_cycles(rows)
    for row in rows:
        row["status"] = "invalid" if row.get("errors") else "valid"
        if row.get("memberMatch") and _mapping_value(row, "memberMatch", "status") != "matched":
            row["status"] = "unmatched"


def _validate_explicit_upstream_id(
    row: dict[str, object], snapshot: Mapping[str, object] | None
) -> None:
    values = _row_values(row)
    entity = str(row.get("entity"))
    candidate_keys = {
        "projects": ("timecueProjectId", "projects"),
        "tasks": ("timecueTaskId", "tasks"),
        "skills": ("timecueSpecialtyId", "specialties"),
    }
    candidate = candidate_keys.get(entity)
    if candidate is None:
        return
    candidate_key, collection_key = candidate
    explicit_id = _string_value(values.get(candidate_key))
    if not explicit_id or snapshot is None:
        return
    collection = snapshot.get(collection_key)
    if not isinstance(collection, list):
        return
    ids = {
        str(item.get("id")) for item in collection if isinstance(item, Mapping) and item.get("id")
    }
    if explicit_id not in ids:
        _append_error(row, f"upstream_id_not_found:{candidate_key}")
    else:
        row["mapping"] = {
            **_mapping(row.get("mapping")),
            "upstreamId": explicit_id,
            "matchType": "explicit",
        }
        row["operation"] = "verify"


def _match_member(row: dict[str, object], member_rows: Sequence[Mapping[str, object]]) -> None:
    requested = _string_value(_row_values(row).get("timecueMemberId"))
    index: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    for member in member_rows:
        for key in ("id", "memberId", "organizationMemberId", "timecueMemberId", "userId"):
            value = _string_value(member.get(key))
            if value:
                index[value].append(member)
    if not requested:
        return
    matches = index.get(requested, [])
    if len(matches) == 1:
        member = matches[0]
        worker_profile_id = _member_worker_profile_id(member)
        row["memberMatch"] = {
            "status": "matched",
            "requestedId": requested,
            "memberId": _member_id(member) or requested,
            "workerProfileId": worker_profile_id,
            "matchedBy": "exact_id",
        }
        if worker_profile_id:
            row["mapping"] = {
                **_mapping(row.get("mapping")),
                "workerProfileId": worker_profile_id,
            }
        return
    status = "ambiguous" if len(matches) > 1 else "unmatched"
    row["memberMatch"] = {"status": status, "requestedId": requested}
    _append_error(row, f"member_{status}:{requested}")


def _member_worker_profile_id(member: Mapping[str, object]) -> str | None:
    for key in ("workerProfileId", "worker_profile_id", "workerId", "profileId"):
        value = _string_value(member.get(key))
        if value:
            return value
    nested = member.get("workerProfile")
    if isinstance(nested, Mapping):
        return _string_value(nested.get("id"))
    return None


def _member_id(member: Mapping[str, object]) -> str | None:
    for key in ("id", "memberId", "organizationMemberId", "timecueMemberId"):
        value = _string_value(member.get(key))
        if value:
            return value
    return None


def _mark_planning_cycles(rows: list[dict[str, object]]) -> None:
    planning = {
        _string_value(_row_values(row).get("taskKey")): row
        for row in rows
        if row.get("entity") == "planning" and _string_value(_row_values(row).get("taskKey"))
    }
    graph = {
        task_key: set(_split_keys(_row_values(row).get("prerequisiteKeys")))
        for task_key, row in planning.items()
        if task_key
    }
    visiting: set[str] = set()
    visited: set[str] = set()
    cycles: set[str] = set()

    def visit(task_key: str, path: tuple[str, ...]) -> None:
        if task_key in visiting:
            cycles.update(path[path.index(task_key) :])
            return
        if task_key in visited:
            return
        visiting.add(task_key)
        for dependency in graph.get(task_key, set()):
            if dependency in graph:
                visit(dependency, (*path, dependency))
        visiting.remove(task_key)
        visited.add(task_key)

    for task_key in graph:
        visit(task_key, (task_key,))
    for task_key in cycles:
        row = planning.get(task_key)
        if row is not None:
            _append_error(row, "planning_dependency_cycle")


def _timecue_project_status(value: str | None) -> str:
    """Map spreadsheet statuses onto the verified Timecue project enum."""

    raw = (value or "scheduled").strip().lower().replace(" ", "_")
    mapped = PROJECT_STATUS_ALIASES.get(raw, raw)
    return mapped if mapped in TIMECUE_PROJECT_STATUSES else "scheduled"


def _project_create_body(values: Mapping[str, object]) -> dict[str, object]:
    """Build the seed-verified project create payload.

    Timecue rejects spreadsheet aliases such as ``planned``. Currency and
    country code are required by the existing create route even when the sheet
    omits them.
    """

    body: dict[str, object] = {
        "name": _string_value(values.get("name")) or "",
        "timezone": _string_value(values.get("timezone")) or "Europe/Warsaw",
        "status": _timecue_project_status(_string_value(values.get("status"))),
        "currency": _string_value(values.get("currency")) or DEFAULT_PROJECT_CURRENCY,
        "countryCode": _string_value(values.get("countryCode")) or DEFAULT_PROJECT_COUNTRY_CODE,
    }
    description = _string_value(values.get("description"))
    if description:
        body["description"] = description
    return body


def _upstream_rejection_message(
    exc: UpstreamIntegrationError, *, after_refresh: bool = False
) -> str:
    prefix = (
        "Timecue rejected the import request after session refresh"
        if after_refresh
        else "Timecue rejected the import request"
    )
    detail = exc.detail.strip() if isinstance(exc.detail, str) and exc.detail.strip() else None
    return f"{prefix}. {detail}" if detail else f"{prefix}."


def _commit_projects(
    rows: list[dict[str, object]],
    project_ids: dict[str, str],
    upstream: RouteClient,
    receipts: list[dict[str, object]],
    receipt_by_row: dict[str, dict[str, object]],
    organization_id: str,
    store: Store,
    record: dict[str, object],
) -> None:
    for row in rows:
        if row.get("entity") != "projects":
            continue
        row_key = str(row.get("rowKey"))
        if _receipt_is_terminal(receipt_by_row.get(row_key)):
            _remember_mapping(row, project_ids)
            continue
        values = _row_values(row)
        external_key = _string_value(values.get("externalKey"))
        mapping = _mapping(row.get("mapping"))
        existing_id = _string_value(mapping.get("upstreamId"))
        if existing_id and external_key:
            result = _read_project(upstream, organization_id, existing_id)
            if result is None:
                receipt = _receipt(
                    row,
                    status="needs_reconciliation",
                    code="project_readback_missing",
                    message="The mapped project was not found during readback.",
                    upstream_id=existing_id,
                )
            else:
                project_ids[external_key] = existing_id
                receipt = _receipt(
                    row,
                    status="verified",
                    message="Existing project mapping verified; no update was sent.",
                    upstream_id=existing_id,
                    readback=result,
                )
            _upsert_receipt(receipts, receipt_by_row, receipt)
            _persist_progress(store, organization_id, record, rows, receipts)
            continue
        body = _project_create_body(values)
        try:
            response = upstream.request(
                "POST", f"/organizations/{organization_id}/projects", json_body=body
            )
            upstream_id = _response_id(response)
            readback = (
                _read_project(upstream, organization_id, upstream_id) if upstream_id else None
            )
            if not upstream_id or readback is None:
                receipt = _receipt(
                    row,
                    status="needs_reconciliation",
                    code="project_readback_missing",
                    message=(
                        "Timecue accepted the project request but canonical readback "
                        "was incomplete."
                    ),
                    upstream_id=upstream_id,
                    response=response,
                )
            else:
                if external_key:
                    project_ids[external_key] = upstream_id
                row["mapping"] = {
                    **mapping,
                    "upstreamId": upstream_id,
                    "matchType": "created",
                }
                receipt = _receipt(
                    row,
                    status="applied",
                    message="Project created and read back from Timecue.",
                    upstream_id=upstream_id,
                    response=response,
                    readback=readback,
                )
        except ImportServiceError as exc:
            receipt = _receipt(row, status="failed", code=exc.code, message=exc.message)
        _upsert_receipt(receipts, receipt_by_row, receipt)
        _persist_progress(store, organization_id, record, rows, receipts)


def _commit_tasks(
    rows: list[dict[str, object]],
    project_ids: Mapping[str, str],
    task_ids: dict[str, str],
    upstream: RouteClient,
    receipts: list[dict[str, object]],
    receipt_by_row: dict[str, dict[str, object]],
    organization_id: str,
    store: Store,
    record: dict[str, object],
) -> None:
    for row in rows:
        if row.get("entity") != "tasks":
            continue
        row_key = str(row.get("rowKey"))
        if _receipt_is_terminal(receipt_by_row.get(row_key)):
            _remember_mapping(row, task_ids)
            continue
        values = _row_values(row)
        external_key = _string_value(values.get("externalKey"))
        project_key = _string_value(values.get("projectKey"))
        project_id = project_ids.get(project_key or "")
        if not project_id:
            receipt = _receipt(
                row,
                status="blocked",
                code="project_dependency_failed",
                message="The task waits for its project mapping to be applied.",
            )
            _upsert_receipt(receipts, receipt_by_row, receipt)
            continue
        mapping = _mapping(row.get("mapping"))
        existing_id = _string_value(mapping.get("upstreamId"))
        if existing_id and external_key:
            try:
                readback = _read_task(upstream, organization_id, project_id, existing_id)
                if readback is None:
                    receipt = _receipt(
                        row,
                        status="needs_reconciliation",
                        code="task_readback_missing",
                        message="The mapped task was not found during readback.",
                        upstream_id=existing_id,
                    )
                else:
                    task_ids[external_key] = existing_id
                    desired_dates = _task_date_values(values)
                    if desired_dates and not _task_dates_match(readback, desired_dates):
                        patch_path = (
                            f"/organizations/{organization_id}/projects/{project_id}/tasks/"
                            f"{existing_id}"
                        )
                        before = {
                            key: readback.get(key) for key in ("plannedStartAt", "plannedEndAt")
                        }
                        response = upstream.request("PATCH", patch_path, json_body=desired_dates)
                        patched = _read_task(upstream, organization_id, project_id, existing_id)
                        if patched is None or not _task_dates_match(patched, desired_dates):
                            receipt = _receipt(
                                row,
                                status="needs_reconciliation",
                                code="task_date_readback_mismatch",
                                message=(
                                    "The task date update returned without the intended "
                                    "canonical readback."
                                ),
                                upstream_id=existing_id,
                                response=response,
                                readback=patched,
                                precondition=before,
                            )
                        else:
                            receipt = _receipt(
                                row,
                                status="applied",
                                message="Existing task dates updated and read back from Timecue.",
                                upstream_id=existing_id,
                                response=response,
                                readback=patched,
                                precondition=before,
                            )
                    else:
                        receipt = _receipt(
                            row,
                            status="verified",
                            message="Existing task mapping and dates verified; no update was sent.",
                            upstream_id=existing_id,
                            readback=readback,
                        )
            except ImportServiceError as exc:
                receipt = _receipt(
                    row,
                    status="failed",
                    code=exc.code,
                    message=exc.message,
                    upstream_id=existing_id,
                )
            _upsert_receipt(receipts, receipt_by_row, receipt)
            _persist_progress(store, organization_id, record, rows, receipts)
            continue
        body: dict[str, object] = {
            "projectId": project_id,
            "title": _string_value(values.get("title")) or "",
            "status": _string_value(values.get("status")) or "planned",
        }
        for source_key, target_key in (
            ("plannedStart", "plannedStartAt"),
            ("plannedEnd", "plannedEndAt"),
            ("description", "description"),
        ):
            value = _string_value(values.get(source_key))
            if value:
                body[target_key] = value
        estimated_minutes = _number_value(values.get("estimatedMinutes"))
        if estimated_minutes is not None:
            body["estimatedMinutes"] = int(estimated_minutes)
        try:
            response = upstream.request(
                "POST",
                f"/organizations/{organization_id}/projects/{project_id}/tasks",
                json_body=body,
            )
            upstream_id = _response_id(response)
            readback = (
                _read_task(upstream, organization_id, project_id, upstream_id)
                if upstream_id
                else None
            )
            if not upstream_id or readback is None:
                receipt = _receipt(
                    row,
                    status="needs_reconciliation",
                    code="task_readback_missing",
                    message=(
                        "Timecue accepted the task request but canonical readback was incomplete."
                    ),
                    upstream_id=upstream_id,
                    response=response,
                )
            elif not _task_dates_match(readback, _task_date_values(values)):
                receipt = _receipt(
                    row,
                    status="needs_reconciliation",
                    code="task_date_readback_mismatch",
                    message="The created task did not read back the intended dates.",
                    upstream_id=upstream_id,
                    response=response,
                    readback=readback,
                )
            else:
                if external_key:
                    task_ids[external_key] = upstream_id
                row["mapping"] = {
                    **mapping,
                    "upstreamId": upstream_id,
                    "matchType": "created",
                }
                receipt = _receipt(
                    row,
                    status="applied",
                    message="Task created and read back from Timecue.",
                    upstream_id=upstream_id,
                    response=response,
                    readback=readback,
                )
        except ImportServiceError as exc:
            receipt = _receipt(row, status="failed", code=exc.code, message=exc.message)
        _upsert_receipt(receipts, receipt_by_row, receipt)
        _persist_progress(store, organization_id, record, rows, receipts)


def _commit_assignments(
    rows: list[dict[str, object]],
    project_ids: Mapping[str, str],
    task_ids: Mapping[str, str],
    worker_profile_ids: dict[str, str],
    upstream: RouteClient,
    receipts: list[dict[str, object]],
    receipt_by_row: dict[str, dict[str, object]],
    organization_id: str,
    store: Store,
    record: dict[str, object],
) -> None:
    by_task: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        if row.get("entity") == "assignments":
            task_key = _string_value(_row_values(row).get("taskKey"))
            if task_key:
                by_task[task_key].append(row)
    for task_key, assignment_rows in by_task.items():
        if all(
            _receipt_is_terminal(receipt_by_row.get(str(row.get("rowKey"))))
            for row in assignment_rows
        ):
            continue
        task_row = next(
            (
                row
                for row in rows
                if row.get("entity") == "tasks"
                and _string_value(_row_values(row).get("externalKey")) == task_key
            ),
            None,
        )
        task_id = task_ids.get(task_key)
        project_key = _string_value(_row_values(task_row).get("projectKey")) if task_row else None
        project_id = project_ids.get(project_key or "")
        if not task_id or not project_id:
            for row in assignment_rows:
                receipt = _receipt(
                    row,
                    status="blocked",
                    code="task_dependency_failed",
                    message="The assignment waits for its task mapping to be applied.",
                )
                _upsert_receipt(receipts, receipt_by_row, receipt)
            continue
        requested_profiles: set[str] = set()
        missing_profile_rows: list[dict[str, object]] = []
        for row in assignment_rows:
            worker_key = _string_value(_row_values(row).get("workerKey"))
            worker_row = next(
                (
                    candidate
                    for candidate in rows
                    if candidate.get("entity") == "workers"
                    and _string_value(_row_values(candidate).get("externalKey")) == worker_key
                ),
                None,
            )
            profile_id = (
                _string_value(_mapping(worker_row.get("mapping")).get("workerProfileId"))
                if worker_row
                else None
            )
            if profile_id:
                worker_profile_ids[worker_key or ""] = profile_id
                requested_profiles.add(profile_id)
            else:
                missing_profile_rows.append(row)
        for row in missing_profile_rows:
            receipt = _receipt(
                row,
                status="manual",
                code="worker_profile_unavailable",
                message="The matched member has no verified worker profile for assignment.",
            )
            _upsert_receipt(receipts, receipt_by_row, receipt)
        if not requested_profiles:
            continue
        assignment_path = (
            f"/organizations/{organization_id}/projects/{project_id}/tasks/{task_id}/assignments"
        )
        try:
            current = upstream.request("GET", assignment_path)
            current_workers = _assignment_worker_ids(current, "directWorkers")
            current_teams = _assignment_team_ids(current)
            desired_workers = sorted(set(current_workers).union(requested_profiles))
            response = upstream.request(
                "PUT",
                assignment_path,
                json_body={"workerProfileIds": desired_workers, "teamIds": current_teams},
            )
            readback = upstream.request("GET", assignment_path)
            effective = set(_assignment_worker_ids(readback, "effectiveWorkers"))
            if not requested_profiles.issubset(effective):
                for row in assignment_rows:
                    if row in missing_profile_rows:
                        continue
                    receipt = _receipt(
                        row,
                        status="needs_reconciliation",
                        code="assignment_readback_missing",
                        message="The assignment write did not read back all requested workers.",
                        upstream_id=task_id,
                        response=response,
                        readback=readback,
                    )
                    _upsert_receipt(receipts, receipt_by_row, receipt)
            else:
                for row in assignment_rows:
                    if row in missing_profile_rows:
                        continue
                    receipt = _receipt(
                        row,
                        status="applied",
                        message=(
                            "Assignment applied with existing direct workers and teams preserved."
                        ),
                        upstream_id=task_id,
                        response=response,
                        readback=readback,
                    )
                    _upsert_receipt(receipts, receipt_by_row, receipt)
        except ImportServiceError as exc:
            for row in assignment_rows:
                if row in missing_profile_rows:
                    continue
                receipt = _receipt(row, status="failed", code=exc.code, message=exc.message)
                _upsert_receipt(receipts, receipt_by_row, receipt)
        _persist_progress(store, organization_id, record, rows, receipts)


def _read_project(
    upstream: RouteClient, organization_id: str, project_id: str
) -> dict[str, object] | None:
    response = upstream.request("GET", f"/organizations/{organization_id}/projects")
    return _find_by_id(response, project_id)


def _read_task(
    upstream: RouteClient, organization_id: str, project_id: str, task_id: str
) -> dict[str, object] | None:
    response = upstream.request(
        "GET",
        f"/organizations/{organization_id}/projects/{project_id}/tasks",
        params={"include_completed": "true"},
    )
    return _find_by_id(response, task_id)


def _task_date_values(values: Mapping[str, object]) -> dict[str, str]:
    result: dict[str, str] = {}
    for source_key, target_key in (
        ("plannedStart", "plannedStartAt"),
        ("plannedEnd", "plannedEndAt"),
    ):
        value = _string_value(values.get(source_key))
        if value:
            result[target_key] = value
    return result


def _task_dates_match(task: Mapping[str, object], desired: Mapping[str, object]) -> bool:
    return all(_datetime_values_equal(task.get(key), value) for key, value in desired.items())


def _datetime_values_equal(left: object, right: object) -> bool:
    left_text = _string_value(left)
    right_text = _string_value(right)
    if left_text is None or right_text is None:
        return left_text == right_text
    try:
        left_value = datetime.fromisoformat(left_text.replace("Z", "+00:00"))
        right_value = datetime.fromisoformat(right_text.replace("Z", "+00:00"))
    except ValueError:
        return left_text == right_text
    if left_value.tzinfo is None or right_value.tzinfo is None:
        return left_text == right_text
    return left_value == right_value


def _find_by_id(response: object, record_id: str | None) -> dict[str, object] | None:
    if not record_id:
        return None
    values = _response_items(response)
    return next(
        (dict(item) for item in values if _string_value(item.get("id")) == record_id),
        None,
    )


def _assignment_worker_ids(response: object, key: str) -> list[str]:
    if not isinstance(response, Mapping):
        return []
    values = response.get(key)
    if not isinstance(values, list):
        return []
    result: list[str] = []
    for value in values:
        if isinstance(value, Mapping):
            worker_id = _string_value(
                value.get("workerProfileId") or value.get("workerId") or value.get("id")
            )
        else:
            worker_id = _string_value(value)
        if worker_id and worker_id not in result:
            result.append(worker_id)
    return result


def _assignment_team_ids(response: object) -> list[str]:
    if not isinstance(response, Mapping):
        return []
    values = response.get("teams")
    if not isinstance(values, list):
        return []
    result: list[str] = []
    for value in values:
        if isinstance(value, Mapping):
            team_id = _string_value(value.get("teamId") or value.get("id"))
        else:
            team_id = _string_value(value)
        if team_id and team_id not in result:
            result.append(team_id)
    return result


def _response_items(response: object) -> list[Mapping[str, object]]:
    if isinstance(response, list):
        return [item for item in response if isinstance(item, Mapping)]
    if isinstance(response, Mapping):
        values = response.get("items")
        if isinstance(values, list):
            return [item for item in values if isinstance(item, Mapping)]
    return []


def _response_id(response: object) -> str | None:
    if isinstance(response, Mapping):
        for key in ("id", "projectId", "taskId", "workerProfileId"):
            value = _string_value(response.get(key))
            if value:
                return value
        nested = response.get("data")
        if isinstance(nested, Mapping):
            return _response_id(nested)
    return None


def _response_payload(response: object) -> object:
    try:
        import httpx

        if isinstance(response, httpx.Response):
            if response.status_code >= 400:
                raise ImportServiceError(
                    "timecue_request_failed",
                    f"Timecue rejected the import request. {_public_error_detail(response)}",
                    status_code=502,
                )
            if response.status_code == 204 or not response.content:
                return {}
            try:
                return response.json()
            except ValueError as exc:
                raise ImportServiceError(
                    "timecue_response_invalid",
                    "Timecue returned a non-JSON import response.",
                    status_code=502,
                ) from exc
    except ImportServiceError:
        raise
    return response


def _source_files(
    uploads: Sequence[UploadedImport], tables: Sequence[ParsedTable]
) -> list[dict[str, object]]:
    entities_by_file: dict[str, list[str]] = defaultdict(list)
    for table in tables:
        if table.entity not in entities_by_file[table.source_file]:
            entities_by_file[table.source_file].append(table.entity)
    return [
        {
            "name": _safe_filename(upload.filename),
            "contentType": upload.content_type,
            "sizeBytes": len(upload.content),
            "sha256": hashlib.sha256(upload.content).hexdigest(),
            "entities": entities_by_file.get(_safe_filename(upload.filename), []),
        }
        for upload in uploads
    ]


def _source_hash(source_files: Sequence[Mapping[str, object]]) -> str:
    canonical = [
        {key: item.get(key) for key in ("name", "sizeBytes", "sha256", "entities")}
        for item in source_files
    ]
    return hashlib.sha256(_canonical_json(canonical).encode("utf-8")).hexdigest()


def _previous_mappings(records: Sequence[Mapping[str, object]]) -> dict[str, dict[str, object]]:
    mappings: dict[str, dict[str, object]] = {}
    for record in records:
        for row_value in _mapping_list(record.get("rows")):
            mapping = _mapping(row_value.get("mapping"))
            if mapping.get("upstreamId") or mapping.get("workerProfileId"):
                key = _mapping_key(row_value)
                mappings[key] = mapping
    return mappings


def _mapping_key(row: Mapping[str, object]) -> str:
    return f"{row.get('entity')}:{row.get('rowKey')}"


def _stable_row_key(entity: str, values: Mapping[str, object], source_row: int) -> str:
    if entity == "assignments":
        first = _string_value(values.get("taskKey")) or f"row-{source_row}"
        second = _string_value(values.get("workerKey")) or f"row-{source_row}"
        return f"{entity}:{_key_part(first)}:{_key_part(second)}"
    if entity == "planning":
        if _is_transfer_row(values):
            parts = (
                _string_value(values.get("fromProjectKey")) or f"row-{source_row}",
                _string_value(values.get("toProjectKey")) or f"row-{source_row}",
                _string_value(values.get("workerKey")) or f"row-{source_row}",
                _string_value(values.get("startsAt")) or f"row-{source_row}",
            )
            return f"{entity}:transfer:{':'.join(_key_part(part) for part in parts)}"
        task_key = _string_value(values.get("taskKey")) or f"row-{source_row}"
        return f"{entity}:task:{_key_part(task_key)}"
    external_key = _string_value(values.get("externalKey")) or f"row-{source_row}"
    return f"{entity}:{_key_part(external_key)}"


def _key_part(value: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip())
    return normalized[:128] or "empty"


def _canonical_headers(raw_headers: Sequence[object]) -> tuple[list[str], list[tuple[str, str]]]:
    headers: list[str] = []
    errors: list[tuple[str, str]] = []
    seen: set[str] = set()
    known = {key: key for columns in TEMPLATE_COLUMNS.values() for key in columns}
    for raw_header in raw_headers:
        text = _string_value(raw_header) or ""
        normalized = _normalize_header(text)
        canonical = next(
            (key for key in known if _normalize_header(key) == normalized),
            text.strip(),
        )
        if not canonical:
            errors.append(("header_empty", "Header names must not be empty."))
            canonical = f"unknown_{len(headers)}"
        if canonical in seen:
            errors.append(("header_duplicate", f"Header {canonical!r} appears more than once."))
        seen.add(canonical)
        headers.append(canonical)
    return headers, errors


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.strip().lower())


def _row_from_values(headers: Sequence[str], values: Sequence[object]) -> dict[str, object]:
    result: dict[str, object] = {}
    for index, header in enumerate(headers):
        value = values[index] if index < len(values) else ""
        result[header] = _json_cell(value)
    return result


def _json_cell(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
            return value[:MAX_CELL_CHARS]
        return value
    return str(value)[:MAX_CELL_CHARS]


def _cell_value(cell: object) -> object:
    return getattr(cell, "value", cell)


def _cell_formula(cell: object) -> bool:
    return str(getattr(cell, "data_type", "")) == "f" or _is_formula(_cell_value(cell))


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.lstrip().startswith("=")


def _entity_from_filename(filename: str) -> str | None:
    stem = Path(filename).stem.lower().strip()
    return ENTITY_ALIASES.get(stem)


def _entity_from_sheet(sheet: str) -> str | None:
    normalized = _normalize_header(sheet)
    for entity, title in ENTITY_SHEETS.items():
        if normalized in {_normalize_header(entity), _normalize_header(title)}:
            return entity
    return None


def _safe_filename(filename: str) -> str:
    basename = Path(filename or "upload").name
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", basename)
    return cleaned[:160] or "upload"


def _operation_for(entity: str) -> str:
    if entity in {"projects", "tasks"}:
        return "create_or_verify"
    if entity == "assignments":
        return "add_or_preserve_assignment"
    if entity == "planning":
        return "proposed_planning_overlay"
    return "local_validation"


def _validate_datetime(value: object, field: str, errors: list[str]) -> None:
    text = _string_value(value)
    if not text:
        return
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        errors.append(f"{field}_invalid")
        return
    if parsed.tzinfo is None:
        errors.append(f"{field}_timezone_required")


def _validate_ordered_datetimes(start: object, end: object, errors: list[str]) -> None:
    start_text = _string_value(start)
    end_text = _string_value(end)
    if not start_text or not end_text:
        return
    try:
        start_at = datetime.fromisoformat(start_text.replace("Z", "+00:00"))
        end_at = datetime.fromisoformat(end_text.replace("Z", "+00:00"))
    except ValueError:
        return
    if start_at.tzinfo is not None and end_at.tzinfo is not None and end_at <= start_at:
        errors.append("planned_end_not_after_start")


def _validate_coordinates(values: Mapping[str, object], errors: list[str]) -> None:
    latitude = _number_value(values.get("defaultLatitude"))
    longitude = _number_value(values.get("defaultLongitude"))
    if (latitude is None) != (longitude is None):
        errors.append("coordinates_pair_required")
        return
    if latitude is not None and not -90 <= latitude <= 90:
        errors.append("latitude_invalid")
    if longitude is not None and not -180 <= longitude <= 180:
        errors.append("longitude_invalid")


def _validate_weather_rules(values: Mapping[str, object], errors: list[str]) -> None:
    raw = _string_value(values.get("weatherRulesJson"))
    if not raw:
        return
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        errors.append("weather_rules_json_invalid")
        return
    if not isinstance(parsed, Mapping):
        errors.append("weather_rules_object_required")
        return
    if parsed.get("confirmed") is not True:
        errors.append("weather_rules_unconfirmed")
    rules = parsed.get("rules")
    if not isinstance(rules, list) or not rules:
        errors.append("weather_rules_required")
    capacity = parsed.get("capacity", parsed.get("defaultCapacity"))
    numeric_capacity = _number_value(capacity)
    if numeric_capacity is None or not 0 <= numeric_capacity <= 1:
        errors.append("weather_capacity_invalid")
    if isinstance(rules, list):
        for index, rule in enumerate(rules):
            if not isinstance(rule, Mapping):
                errors.append(f"weather_rule_invalid:{index}")
                continue
            metric = _string_value(rule.get("metric"))
            operator = _string_value(rule.get("operator"))
            threshold = _number_value(rule.get("threshold"))
            if not metric or not operator or threshold is None:
                errors.append(f"weather_rule_invalid:{index}")


def _validate_planning_values(values: Mapping[str, object], errors: list[str]) -> None:
    numeric_fields = (
        "optimisticPersonHours",
        "mostLikelyPersonHours",
        "pessimisticPersonHours",
        "minCrew",
        "maxCrew",
        "priority",
    )
    numeric: dict[str, float] = {}
    for field in numeric_fields:
        value = _number_value(values.get(field))
        if values.get(field) not in (None, "") and value is None:
            errors.append(f"{field}_invalid")
        elif value is not None:
            numeric[field] = value
            if value < 0:
                errors.append(f"{field}_negative")
    if numeric.get("minCrew", 0) > numeric.get("maxCrew", numeric.get("minCrew", 0)):
        errors.append("crew_bounds_invalid")
    effort = [numeric.get(field) for field in numeric_fields[:3]]
    if all(value is not None for value in effort) and not (effort[0] <= effort[1] <= effort[2]):
        errors.append("effort_range_invalid")


def _require_reference(row: dict[str, object], field: str, value: object, known: set[str]) -> None:
    text = _string_value(value)
    if text and text not in known:
        _append_error(row, f"reference_not_found:{field}:{text}")


def _append_error(row: dict[str, object], message: str) -> None:
    errors = _string_list(row.get("errors"))
    if message not in errors:
        errors.append(message)
    row["errors"] = errors


def _row_errors(rows: Sequence[Mapping[str, object]]) -> list[dict[str, object]]:
    result: list[dict[str, object]] = []
    for row in rows:
        for error in _string_list(row.get("errors")):
            result.append(
                {
                    "code": error.split(":", 1)[0],
                    "message": error,
                    "rowKey": row.get("rowKey"),
                    "entity": row.get("entity"),
                    "sourceFile": row.get("sourceFile"),
                    "sourceRow": row.get("sourceRow"),
                }
            )
    return result


def _counts(
    rows: Sequence[Mapping[str, object]], errors: Sequence[Mapping[str, object]]
) -> dict[str, object]:
    per_entity: dict[str, dict[str, int]] = {}
    for entity in ENTITY_ORDER:
        selected = [row for row in rows if row.get("entity") == entity]
        per_entity[entity] = {
            "total": len(selected),
            "valid": sum(1 for row in selected if row.get("status") == "valid"),
            "invalid": sum(1 for row in selected if row.get("status") != "valid"),
            "creates": sum(
                1
                for row in selected
                if row.get("status") == "valid" and row.get("operation") == "create_or_verify"
            ),
        }
    return {
        "total": len(rows),
        "valid": sum(1 for row in rows if row.get("status") == "valid"),
        "invalid": sum(1 for row in rows if row.get("status") != "valid"),
        "errors": len(errors),
        "unmatchedMembers": sum(
            1
            for row in rows
            if _mapping_value(row, "memberMatch", "status") in {"unmatched", "ambiguous"}
        ),
        **per_entity,
    }


def _record_hash(record: Mapping[str, object]) -> str:
    ignored = {"id", "createdAt", "updatedAt", "previewHash"}
    value = {key: item for key, item in record.items() if key not in ignored}
    return hashlib.sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _canonical_json(value: object) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)


def _seed_existing_mappings(
    rows: Sequence[Mapping[str, object]],
    project_ids: dict[str, str],
    task_ids: dict[str, str],
    worker_profile_ids: dict[str, str],
) -> None:
    for row in rows:
        values = _row_values(row)
        external_key = _string_value(values.get("externalKey"))
        mapping = _mapping(row.get("mapping"))
        if not external_key:
            continue
        upstream_id = _string_value(mapping.get("upstreamId"))
        profile_id = _string_value(mapping.get("workerProfileId"))
        if row.get("entity") == "projects" and upstream_id:
            project_ids[external_key] = upstream_id
        elif row.get("entity") == "tasks" and upstream_id:
            task_ids[external_key] = upstream_id
        elif row.get("entity") == "workers" and profile_id:
            worker_profile_ids[external_key] = profile_id


def _remember_mapping(row: Mapping[str, object], target: dict[str, str]) -> None:
    values = _row_values(row)
    external_key = _string_value(values.get("externalKey"))
    upstream_id = _string_value(_mapping(row.get("mapping")).get("upstreamId"))
    if external_key and upstream_id:
        target[external_key] = upstream_id


def _receipt(
    row: Mapping[str, object],
    *,
    status: str,
    message: str,
    code: str | None = None,
    upstream_id: str | None = None,
    response: object | None = None,
    readback: object | None = None,
    precondition: object | None = None,
) -> dict[str, object]:
    value: dict[str, object] = {
        "id": str(uuid4()),
        "rowKey": row.get("rowKey"),
        "entity": row.get("entity"),
        "operation": row.get("operation"),
        "status": status,
        "message": message,
        "recordedAt": _utc_now(),
    }
    if code:
        value["code"] = code
    if upstream_id:
        value["upstreamId"] = upstream_id
    if response is not None:
        value["response"] = _redact_response(response)
    if readback is not None:
        value["readback"] = _redact_response(readback)
    if precondition is not None:
        value["precondition"] = _redact_response(precondition)
    return value


def _redact_response(value: object) -> object:
    if isinstance(value, Mapping):
        return {
            str(key): _redact_response(item)
            for key, item in value.items()
            if str(key).lower()
            not in {
                "access_token",
                "refresh_token",
                "authorization",
                "password",
                "secret",
                "api_key",
            }
        }
    if isinstance(value, list):
        return [_redact_response(item) for item in value[:256]]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)[:MAX_CELL_CHARS]


def _upsert_receipt(
    receipts: list[dict[str, object]],
    receipt_by_row: dict[str, dict[str, object]],
    receipt: dict[str, object],
) -> None:
    row_key = str(receipt.get("rowKey"))
    prior = receipt_by_row.get(row_key)
    if prior is None:
        receipts.append(receipt)
    else:
        index = receipts.index(prior)
        receipts[index] = receipt
    receipt_by_row[row_key] = receipt


def _receipt_is_terminal(receipt: Mapping[str, object] | None) -> bool:
    return bool(receipt and receipt.get("status") in {"applied", "verified", "manual"})


def _persist_progress(
    store: Store,
    organization_id: str,
    record: dict[str, object],
    rows: list[dict[str, object]],
    receipts: list[dict[str, object]],
) -> None:
    record["rows"] = rows
    record["receipts"] = receipts
    store.save_record(organization_id, IMPORT_RECORD_KIND, str(record["id"]), record)


def _batch_status(receipts: Sequence[Mapping[str, object]]) -> str:
    statuses = {str(receipt.get("status")) for receipt in receipts}
    if statuses.intersection({"failed", "blocked", "manual", "needs_reconciliation"}):
        return (
            "partially_applied"
            if statuses.intersection({"applied", "verified"})
            else "needs_reconciliation"
        )
    return "applied"


def _batch_message(status: object) -> str:
    if status == "applied":
        return "Supported Timecue changes were applied and read back."
    return "The import is partially applied; review receipts and reconcile the remaining steps."


def _row_values(row: Mapping[str, object] | None) -> dict[str, object]:
    if not row:
        return {}
    value = row.get("values")
    return dict(value) if isinstance(value, Mapping) else {}


def _assumption_value(
    assumptions: Mapping[str, object], task_values: Mapping[str, object], field: str
) -> object:
    value = assumptions.get(field)
    return value if value not in (None, "") else task_values.get(field)


def _mapping(value: object) -> dict[str, object]:
    return dict(value) if isinstance(value, Mapping) else {}


def _mapping_list(value: object) -> list[dict[str, object]]:
    if not isinstance(value, list):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]


def _mutable_mapping(value: object) -> dict[str, object]:
    return dict(value) if isinstance(value, Mapping) else {}


def _string_list(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item) for item in value if isinstance(item, (str, int, float))]


def _split_keys(value: object) -> list[str]:
    text = _string_value(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[,;|\n]", text) if part.strip()]


def _number_value(value: object) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _string_value(value: object) -> str | None:
    if isinstance(value, str):
        text = value.strip()
        return text[:MAX_CELL_CHARS] if text else None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)
    return None


def _mapping_value(row: Mapping[str, object], mapping_key: str, value_key: str) -> str | None:
    return _string_value(_mapping(row.get(mapping_key)).get(value_key))


def _utc_now() -> str:
    from datetime import UTC

    return datetime.now(UTC).isoformat()


def _column_letter(number: int) -> str:
    result = ""
    current = number
    while current:
        current, remainder = divmod(current - 1, 26)
        result = chr(65 + remainder) + result
    return result


__all__ = [
    "IMPORT_COMMIT_LEASE",
    "IMPORT_RECORD_KIND",
    "MAX_IMPORT_BYTES",
    "MAX_IMPORT_FILES",
    "MAX_IMPORT_ROWS",
    "ImportServiceError",
    "RouteClient",
    "TimecueRouteClient",
    "UploadedImport",
    "commit_import",
    "merge_confirmed_planning",
    "preview_import",
    "template_bytes",
]
